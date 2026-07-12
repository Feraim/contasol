"""Exportación genérica de tablas de informes a PDF y Excel."""

from io import BytesIO

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COLOR_MARCA = colors.HexColor("#e8671b")
COLOR_FILA_PAR = colors.HexColor("#f6f2ee")


def formatear(valor) -> str:
    if isinstance(valor, bool) or valor is None:
        return "" if valor is None else str(valor)
    if isinstance(valor, (int, float)):
        return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return str(valor)


def a_pdf(titulo: str, columnas: list[str], filas: list[list], subtitulo: str = "") -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()
    elementos = [Paragraph(titulo, estilos["Title"])]
    if subtitulo:
        elementos.append(Paragraph(subtitulo, estilos["Normal"]))
    elementos.append(Spacer(1, 12))

    datos = [columnas] + [[formatear(c) for c in fila] for fila in filas]
    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), COLOR_MARCA),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, COLOR_FILA_PAR]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elementos.append(tabla)
    doc.build(elementos)
    return buffer.getvalue()


def a_excel(titulo: str, columnas: list[str], filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (titulo or "Hoja1")[:31]
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append(list(fila))
    for i, columna in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(columna)) + 2)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


MEDIA_EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def respuesta(
    formato: str, nombre_archivo: str, titulo: str, columnas: list[str], filas: list[list],
    subtitulo: str = "",
) -> Response:
    """Construye la respuesta HTTP de descarga para 'pdf' o 'excel'."""
    if formato == "pdf":
        return Response(
            a_pdf(titulo, columnas, filas, subtitulo),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}.pdf"'},
        )
    return Response(
        a_excel(titulo, columnas, filas),
        media_type=MEDIA_EXCEL,
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}.xlsx"'},
    )
